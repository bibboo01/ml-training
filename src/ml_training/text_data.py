"""Text and document datasets for the HuggingFace backend (Thai and any other language).

Tasks (``data.task: auto`` works it out from the files):
  classify  one label per text / document   (sentiment, intent, document type, spam ...)
  ner       a tag per word -> pull out fields (names, dates, amounts on an invoice ...)

Accepted inputs under ``data/raw/<name>/`` (or ``data.path`` = a folder or one file):

1. A table – .csv .tsv .xlsx .json .jsonl .parquet – with a text column and a label column.
   Names are found automatically (text / sentence / content / review ... and
   label / category / class / intent / sentiment ...) or set with data.text_column /
   data.label_column. A ``split`` column (train / val / test) is respected.
2. NER: JSONL rows ``{"tokens": [...], "tags": [...]}`` (tags = B-PER, I-PER, O ...) or a CoNLL
   file (``word TAG`` per line, blank line between sentences; .conll .bio .iob or .txt).
3. Documents: one sub-folder per class with .txt .md .pdf .docx or images (.jpg .png .tif ...);
   optionally inside train/ val|valid / test/. PDFs use their text layer; scanned PDFs and
   images are read with OCR (EasyOCR, ``uv sync --extra ocr``; languages data.ocr_langs).

``prepare`` copies the raw files to ``data/raw/<name>`` (immutable) and extracts document text
to ``data/interim/<name>/texts/`` (cached). ``featurize`` creates missing val / test splits
(stratified, reproducible) and writes ``data/processed/<name>/{train,val,test}.jsonl`` +
``meta.json``.
"""

from __future__ import annotations

import json
import random
import re
import shutil
from collections import Counter
from pathlib import Path
from typing import Any

import pandas as pd
from omegaconf import DictConfig, OmegaConf

from ml_training.utils import ensure_dir, log, read_json, write_json

TABLE_EXTS = {".csv", ".tsv", ".xlsx", ".xls", ".json", ".jsonl", ".parquet"}
CONLL_EXTS = {".conll", ".bio", ".iob", ".iob2"}
TEXT_EXTS = {".txt", ".md"}
DOC_EXTS = TEXT_EXTS | {".pdf", ".docx"}
IMG_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".tif", ".tiff"}
SPLIT_ALIASES = {
    "train": "train",
    "val": "val",
    "valid": "val",
    "validation": "val",
    "dev": "val",
    "test": "test",
}
TEXT_NAMES = (
    "text",
    "sentence",
    "content",
    "review",
    "message",
    "document",
    "body",
    "comment",
    "question",
    "utterance",
    "title",
    "ข้อความ",
)
LABEL_NAMES = (
    "label",
    "labels",
    "category",
    "class",
    "target",
    "intent",
    "sentiment",
    "topic",
    "tag",
    "y",
    "หมวด",
    "ประเภท",
)
TASKS = ("classify", "ner")


# --------------------------------------------------------------------------- locate
def raw_root(cfg: DictConfig) -> Path:
    return Path(cfg.paths.raw) / cfg.data.name


def _source(cfg: DictConfig) -> Path:
    sample = cfg.data.get("sample")
    if sample and not raw_root(cfg).exists():  # built-in Thai sample (text_samples.py)
        from ml_training.text_samples import make

        return make(sample, raw_root(cfg))
    p = cfg.data.get("path")
    src = Path(p) if p else raw_root(cfg)
    if not src.is_absolute():
        src = Path(cfg.paths.root) / src
    if not src.exists():
        raise FileNotFoundError(
            f"Dataset not found: {src}\nPut a CSV/Excel/JSONL file or a folder of documents in "
            f"data/raw/{cfg.data.name}/ (see README: text and documents)"
        )
    return src


def _input_file(root: Path) -> Path | None:
    """The single data file of a dataset folder (a table or CoNLL), else None (= documents)."""
    if root.is_file():
        return root
    files = [
        p for p in root.iterdir() if p.is_file() and p.suffix.lower() in TABLE_EXTS | CONLL_EXTS
    ]
    subdirs = [p for p in root.iterdir() if p.is_dir() and not p.name.startswith(".")]
    if len(files) == 1 and not subdirs:
        return files[0]
    if files and not subdirs:  # several tables -> train/val/test files or the biggest one
        named = [f for f in files if f.stem.lower() in SPLIT_ALIASES]
        return root if named else max(files, key=lambda f: f.stat().st_size)
    return None


# --------------------------------------------------------------------------- readers
def read_table(path: Path) -> pd.DataFrame:
    ext = path.suffix.lower()
    if ext in (".csv", ".tsv"):
        return pd.read_csv(
            path, sep="\t" if ext == ".tsv" else None, engine="python", encoding="utf-8-sig"
        )
    if ext in (".xlsx", ".xls"):
        return pd.read_excel(path)
    if ext == ".jsonl":
        return pd.read_json(path, lines=True)
    if ext == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, dict):  # {"data": [...]} style
            data = next((v for v in data.values() if isinstance(v, list)), [data])
        return pd.DataFrame(data)
    if ext == ".parquet":
        return pd.read_parquet(path)
    raise ValueError(f"Unsupported table {path}")


def read_conll(path: Path) -> list[dict[str, Any]]:
    rows, tokens, tags = [], [], []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("-DOCSTART-"):
            if tokens:
                rows.append({"tokens": tokens, "tags": tags})
                tokens, tags = [], []
            continue
        parts = line.split()
        tokens.append(parts[0])
        tags.append(parts[-1] if len(parts) > 1 else "O")
    if tokens:
        rows.append({"tokens": tokens, "tags": tags})
    return rows


def _is_conll(path: Path) -> bool:
    if path.suffix.lower() in CONLL_EXTS:
        return True
    if path.suffix.lower() != ".txt":
        return False
    lines = [x for x in path.read_text(encoding="utf-8").splitlines()[:200] if x.strip()]
    tag = re.compile(r"^(O|[BIESU]-\S+)$")
    return bool(lines) and sum(bool(tag.match(x.split()[-1])) for x in lines) >= 0.9 * len(lines)


def find_column(df: pd.DataFrame, wanted: str | None, names: tuple[str, ...], kind: str) -> str:
    if wanted and wanted != "auto":
        if wanted not in df:
            raise KeyError(f"{kind} column {wanted!r} not in {list(df.columns)}")
        return wanted
    lower = {str(c).lower().strip(): c for c in df.columns}
    for n in names:
        if n in lower:
            return lower[n]
    obj = [c for c in df.columns if df[c].dtype == object]
    if kind == "text" and obj:  # longest strings on average
        return max(obj, key=lambda c: df[c].astype(str).str.len().mean())
    if kind == "label":  # fewest distinct values among the rest
        cands = [c for c in df.columns if df[c].nunique() <= max(50, len(df) // 5)]
        if cands:
            return min(cands, key=lambda c: df[c].nunique())
    raise KeyError(f"Cannot find the {kind} column in {list(df.columns)} – set data.{kind}_column")


def _split_name(value: Any) -> str | None:
    return SPLIT_ALIASES.get(str(value).strip().lower())


def table_records(df: pd.DataFrame, cfg: DictConfig) -> tuple[str, list[dict[str, Any]]]:
    """DataFrame -> (task, records). NER if it has token / tag list columns."""
    d = cfg.data
    lower = {str(c).lower(): c for c in df.columns}
    tok = next((lower[k] for k in ("tokens", "words") if k in lower), None)
    tag = next((lower[k] for k in ("tags", "ner_tags", "labels", "bio") if k in lower), None)
    split_col = next((lower[k] for k in ("split", "subset", "set") if k in lower), None)
    if tok and tag and isinstance(df[tok].iloc[0], list | tuple):
        recs = [
            {"tokens": [str(t) for t in r[tok]], "tags": [str(t) for t in r[tag]]}
            for _, r in df.iterrows()
        ]
        task = "ner"
    else:
        tcol = find_column(df, d.get("text_column"), TEXT_NAMES, "text")
        lcol = find_column(df.drop(columns=[tcol]), d.get("label_column"), LABEL_NAMES, "label")
        df = df.dropna(subset=[tcol, lcol])
        recs = [{"text": str(r[tcol]), "label": str(r[lcol]).strip()} for _, r in df.iterrows()]
        log.info("Text column %r, label column %r", tcol, lcol)
        task = "classify"
    if split_col:
        for rec, s in zip(recs, df[split_col], strict=False):
            if _split_name(s):
                rec["split"] = _split_name(s)
    return task, recs


# --------------------------------------------------------------------------- documents
def extract_text(path: Path, ocr_langs: list[str] | None = None) -> str:
    """Plain text of a .txt / .md / .pdf / .docx / image file ('' when nothing readable)."""
    ext = path.suffix.lower()
    if ext in TEXT_EXTS:
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext == ".docx":
        import docx

        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            parts += [" | ".join(c.text for c in row.cells) for row in table.rows]
        return "\n".join(p for p in parts if p.strip())
    if ext == ".pdf":
        from pypdf import PdfReader

        text = "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
        if len(text.strip()) >= 20:
            return text
        return _ocr_pdf(path, ocr_langs) or text  # scanned PDF
    if ext in IMG_EXTS:
        return ocr_image(path, ocr_langs)
    return ""


_OCR: dict[tuple, Any] = {}


def _reader(langs: list[str] | None):
    key = tuple(langs or ["th", "en"])
    if key not in _OCR:
        try:
            import easyocr
        except ImportError as e:
            raise ImportError(
                "Images / scanned PDFs need OCR: .\\run.ps1 setup-ocr (or uv sync --extra ocr)"
            ) from e
        import torch

        _OCR[key] = easyocr.Reader(list(key), gpu=torch.cuda.is_available(), verbose=False)
    return _OCR[key]


def ocr_image(path: Path | Any, langs: list[str] | None = None) -> str:
    lines = _reader(langs).readtext(
        str(path) if isinstance(path, Path) else path, detail=0, paragraph=True
    )
    return "\n".join(lines)


def _ocr_pdf(path: Path, langs: list[str] | None) -> str:
    try:
        import pypdfium2 as pdfium
    except ImportError:
        log.warning("%s has no text layer; install OCR (setup-ocr) to read scanned PDFs", path.name)
        return ""
    import numpy as np

    pdf = pdfium.PdfDocument(str(path))
    return "\n".join(ocr_image(np.array(page.render(scale=2).to_pil()), langs) for page in pdf)


def document_records(root: Path, cache: Path, ocr_langs: list[str] | None) -> list[dict[str, Any]]:
    """<class>/<file> or <split>/<class>/<file> -> records with the extracted text (cached)."""
    recs = []
    exts = DOC_EXTS | IMG_EXTS
    files = sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in exts)
    for f in files:
        rel = f.relative_to(root)
        parts = rel.parts[:-1]
        if not parts:
            continue  # file without a class folder
        split = _split_name(parts[0]) if len(parts) > 1 else None
        label = parts[1] if split else parts[0]
        cached = cache / rel.with_suffix(rel.suffix + ".txt")
        if cached.exists() and cached.stat().st_mtime >= f.stat().st_mtime:
            text = cached.read_text(encoding="utf-8")
        else:
            try:
                text = extract_text(f, ocr_langs)
            except ImportError:
                raise
            except Exception as e:  # noqa: BLE001 - one bad file must not stop the dataset
                log.warning("Skipping %s: %s", rel, e)
                continue
            cached.parent.mkdir(parents=True, exist_ok=True)
            cached.write_text(text, encoding="utf-8")
        if not text.strip():
            log.warning("No text in %s – skipped", rel)
            continue
        rec = {"text": text, "label": label, "file": str(rel)}
        if split:
            rec["split"] = split
        recs.append(rec)
    return recs


# --------------------------------------------------------------------------- pipeline
def load_records(cfg: DictConfig) -> tuple[str, list[dict[str, Any]]]:
    root = raw_root(cfg) if raw_root(cfg).exists() else _source(cfg)
    f = _input_file(root)
    langs = list(cfg.data.get("ocr_langs") or ["th", "en"])
    if f is not None and f.is_dir():  # train.csv / val.csv / test.csv
        task, recs = None, []
        for t in sorted(p for p in f.iterdir() if p.suffix.lower() in TABLE_EXTS | CONLL_EXTS):
            t_task, t_recs = _file_records(t, cfg)
            task = task or t_task
            for r in t_recs:
                r.setdefault("split", _split_name(t.stem) or "train")
            recs += t_recs
        return task or "classify", recs
    if f is not None:
        return _file_records(f, cfg)
    cache = Path(cfg.paths.root) / "data" / "interim" / cfg.data.name / "texts"
    return "classify", document_records(root, cache, langs)


def _file_records(f: Path, cfg: DictConfig) -> tuple[str, list[dict[str, Any]]]:
    if _is_conll(f):
        return "ner", read_conll(f)
    return table_records(read_table(f), cfg)


def prepare(cfg: DictConfig) -> Path:
    src = _source(cfg)
    dst = raw_root(cfg)
    if src.resolve() != dst.resolve() and not (dst.exists() and src.is_relative_to(dst)):
        log.info("Copying %s -> %s (raw data is kept immutable under data/raw)", src, dst)
        if src.is_file():
            ensure_dir(dst)
            shutil.copy2(src, dst / src.name)
        else:
            shutil.copytree(src, dst, dirs_exist_ok=True)
    task, recs = load_records(cfg)
    log.info(
        "Raw text dataset %s: %d %s", dst, len(recs), "sentences" if task == "ner" else "texts"
    )
    return dst


def resolve_task(cfg: DictConfig, detected: str) -> str:
    want = str(cfg.data.get("task") or "auto")
    if want == "auto":
        return detected
    if want not in TASKS:
        raise ValueError(f"data.task must be auto, classify or ner – got {want!r}")
    return want


def split_records(
    recs: list[dict[str, Any]], val_size: float, test_size: float, seed: int, stratify: bool = True
) -> dict[str, list[dict[str, Any]]]:
    """Keep given splits; carve missing val / test out of train, per label when possible."""
    out: dict[str, list] = {"train": [], "val": [], "test": []}
    for r in recs:
        out[r.get("split") or "train"].append(r)
    rng = random.Random(seed)
    key = (lambda r: r.get("label", "")) if stratify else (lambda r: "")
    orig = Counter(key(r) for r in out["train"])  # sizes are shares of the whole dataset
    total = len(out["train"])
    for split, size in (("test", test_size), ("val", val_size)):
        if out[split] or size <= 0 or len(out["train"]) < 4:
            continue
        groups: dict[str, list] = {}
        for r in out["train"]:
            groups.setdefault(key(r), []).append(r)
        take, keep = [], []
        for k, items in groups.items():
            items = items[:]
            rng.shuffle(items)
            n = min(len(items) - 1, round(orig[k] * size)) if orig[k] >= 3 else 0
            take += items[:n]
            keep += items[n:]
        if not take:  # tiny classes: plain random split
            items = out["train"][:]
            rng.shuffle(items)
            n = max(1, min(len(items) - 1, round(total * size)))
            take, keep = items[:n], items[n:]
        out[split], out["train"] = take, keep
    return {k: v for k, v in out.items() if v}


def ner_labels(recs: list[dict[str, Any]]) -> list[str]:
    """BIO tag list with O first and B-/I- next to each other."""
    tags = {t for r in recs for t in r["tags"]}
    ents = sorted({t.split("-", 1)[1] for t in tags if "-" in t})
    out = ["O"] + [f"{p}-{e}" for e in ents for p in ("B", "I") if f"{p}-{e}" in tags]
    return out + sorted(tags - set(out))


def featurize(cfg: DictConfig) -> Path:
    d = cfg.data
    detected, recs = load_records(cfg)
    task = resolve_task(cfg, detected)
    if not recs:
        raise ValueError(f"No usable rows / documents found for {d.name}")
    splits = split_records(
        recs,
        float(d.get("val_size", 0.15)),
        float(d.get("test_size", 0.15)),
        int(cfg.seed),
        bool(d.get("stratify", True)) and task == "classify",
    )
    out = ensure_dir(cfg.paths.processed)
    for old in out.glob("*.jsonl"):
        old.unlink()
    for split, rows in splits.items():
        with (out / f"{split}.jsonl").open("w", encoding="utf-8") as fh:
            for r in rows:
                fh.write(
                    json.dumps({k: v for k, v in r.items() if k != "split"}, ensure_ascii=False)
                    + "\n"
                )
    if task == "ner":
        class_names = ner_labels(recs)
        entities = sorted({t.split("-", 1)[1] for t in class_names if "-" in t})
        counts = {
            s: dict(
                Counter(
                    t.split("-", 1)[1]
                    for r in rows
                    for t in r["tags"]
                    if t.startswith(("B-", "S-", "U-"))
                )
            )
            for s, rows in splits.items()
        }
        lengths = [len(r["tokens"]) for r in recs]
    else:
        names = sorted({r["label"] for r in recs})
        class_names = list(d.get("class_names") or names)
        entities = []
        counts = {s: dict(Counter(r["label"] for r in rows)) for s, rows in splits.items()}
        lengths = [len(r["text"]) for r in recs]
    meta = {
        "dataset": d.name,
        "kind": "text",
        "task": task,
        "class_names": class_names,
        "entities": entities,
        "n_classes": len(class_names),
        "n_rows": {s: len(rows) for s, rows in splits.items()},
        "counts": counts,
        "length": {
            "unit": "words" if task == "ner" else "characters",
            "mean": float(sum(lengths) / len(lengths)),
            "max": int(max(lengths)),
        },
        "data_config": OmegaConf.to_container(d, resolve=True),
    }
    write_json(meta, out / "meta.json")
    log.info("Task: %s  classes: %d  rows: %s", task, len(class_names), meta["n_rows"])
    for s, c in counts.items():
        log.info("  %-5s %s", s, dict(sorted(c.items())))
    return out


def load_split(processed: Path, split: str) -> list[dict[str, Any]]:
    path = Path(processed) / f"{split}.jsonl"
    if not path.exists():
        return []
    return [json.loads(x) for x in path.read_text(encoding="utf-8").splitlines() if x.strip()]


def load_meta(processed: Path) -> dict[str, Any]:
    return read_json(Path(processed) / "meta.json")
